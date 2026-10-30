"""TXT, Markdown, HTML, DOCX and RTF adapters."""
from __future__ import annotations

import re
import zipfile
from pathlib import Path
from typing import Iterator

from bs4 import BeautifulSoup

from .base import (EncryptedError, Metadata, ParsedDocument, ParseError, Segment, TocEntry,
                   chunk_text, extract_year, normalize_text)
from .epub import check_zip_safety, html_to_text

MAX_TEXT_BYTES = 200 * 1024 * 1024


def read_text_file(path: Path) -> str:
    try:
        size = path.stat().st_size
        if size > MAX_TEXT_BYTES:
            raise ParseError("Text file is too large to index (over 200 MB).")
        raw = path.read_bytes()
    except OSError as exc:
        raise ParseError(f"Cannot read file: {exc}") from exc
    if raw[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return raw.decode("utf-16", errors="replace")
    for enc in ("utf-8-sig", "cp1250", "cp1252"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _segments(text: str, label_fn=None) -> Iterator[Segment]:
    for n, (off, piece) in enumerate(chunk_text(text)):
        yield Segment("text", 0, piece, "", off)


def _title_from(text: str, path: Path) -> str:
    return ""


class TxtParser:
    def parse(self, path: Path) -> ParsedDocument:
        text = normalize_text(read_text_file(path))
        return ParsedDocument(Metadata(), _segments(text))


class MarkdownParser:
    def parse(self, path: Path) -> ParsedDocument:
        raw = read_text_file(path)
        meta = Metadata()
        m = re.match(r"^---\s*\n(.*?)\n---\s*\n", raw, re.S)
        if m:
            for line in m.group(1).splitlines():
                k, _, v = line.partition(":")
                k, v = k.strip().lower(), v.strip().strip("\"'")
                if k == "title":
                    meta.title = v
                elif k in ("author", "authors"):
                    meta.authors = [a.strip() for a in re.split(r"[;,]", v.strip("[]")) if a.strip()]
                elif k in ("date", "year"):
                    meta.year = extract_year(v)
                elif k in ("lang", "language"):
                    meta.language = v.split("-")[0].lower()
            raw = raw[m.end():]
        toc = []
        for i, line in enumerate(raw.splitlines()):
            h = re.match(r"^(#{1,6})\s+(.*\S)", line)
            if h:
                if not meta.title and len(h.group(1)) == 1:
                    meta.title = h.group(2).strip()
                toc.append(TocEntry(h.group(2).strip(), 0, len(h.group(1))))
        text = re.sub(r"```.*?```", lambda m: m.group(0).strip("`"), raw, flags=re.S)
        text = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", text)
        text = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", text)
        text = re.sub(r"^#{1,6}\s*", "", text, flags=re.M)
        text = re.sub(r"[*_`>]+", "", text)
        return ParsedDocument(meta, _segments(normalize_text(text)), toc=toc)


class HtmlParser:
    def parse(self, path: Path) -> ParsedDocument:
        raw = read_text_file(path)
        soup = BeautifulSoup(raw, "lxml")
        meta = Metadata()
        if soup.title and soup.title.string:
            meta.title = " ".join(soup.title.string.split())
        for tag in soup.find_all("meta"):
            n = (tag.get("name") or "").lower()
            c = (tag.get("content") or "").strip()
            if n == "author" and c:
                meta.authors = [c]
            elif n == "description" and c:
                meta.description = c
            elif n == "keywords" and c:
                meta.subjects = [k.strip() for k in c.split(",") if k.strip()][:20]
        text, heading = html_to_text(raw)
        if not meta.title:
            meta.title = heading
        return ParsedDocument(meta, _segments(text))


class DocxParser:
    def parse(self, path: Path) -> ParsedDocument:
        try:
            with zipfile.ZipFile(str(path)) as zf:
                check_zip_safety(zf)
                if "EncryptedPackage" in zf.namelist():
                    raise EncryptedError()
        except zipfile.BadZipFile as exc:
            # password-protected DOCX is an OLE container, not a zip
            if path.read_bytes()[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
                raise EncryptedError("Encrypted or legacy Word file (.doc is not supported).") from exc
            raise ParseError(f"Not a valid DOCX file: {exc}") from exc
        import docx
        try:
            d = docx.Document(str(path))
        except Exception as exc:
            raise ParseError(f"Damaged DOCX file: {exc}") from exc
        cp = d.core_properties
        meta = Metadata(title=(cp.title or "").strip(),
                        authors=[a.strip() for a in re.split(r"[;&]", cp.author or "") if a.strip()],
                        description=(cp.subject or "").strip(),
                        language=(cp.language or "").split("-")[0].lower())
        paras, toc = [], []
        for p in d.paragraphs:
            t = p.text.strip()
            if not t:
                continue
            paras.append(t)
            if p.style is not None and p.style.name.lower().startswith(("heading", "title")):
                toc.append(TocEntry(t, 0, 1))
                if not meta.title and p.style.name.lower() == "title":
                    meta.title = t
        for table in d.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    paras.append(" | ".join(cells))
        return ParsedDocument(meta, _segments(normalize_text("\n\n".join(paras))), toc=toc)


class RtfParser:
    def parse(self, path: Path) -> ParsedDocument:
        from striprtf.striprtf import rtf_to_text
        raw = read_text_file(path)
        if not raw.lstrip().startswith("{\\rtf"):
            raise ParseError("Not a valid RTF file.")
        try:
            text = rtf_to_text(raw, errors="ignore")
        except Exception as exc:
            raise ParseError(f"Damaged RTF file: {exc}") from exc
        return ParsedDocument(Metadata(), _segments(normalize_text(text)))
