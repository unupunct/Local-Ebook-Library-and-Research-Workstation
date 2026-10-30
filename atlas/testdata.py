"""Synthetic document generators used by tests, the benchmark and the demo script.

Every document contains known distinctive words so search results can be
checked against exact expectations.
"""
from __future__ import annotations

import random
import zipfile
from pathlib import Path

import pymupdf

LOREM = ("soil water garden plant growth season weather bench timber joint glue finish battery cell charge panel "
         "wind turbine heat storage history story reference chapter section").split()


def filler(rng: random.Random, n: int) -> str:
    return " ".join(rng.choice(LOREM) for _ in range(n))


def make_pdf(path: Path, pages: list[str], title: str = "", author: str = "") -> Path:
    doc = pymupdf.open()
    for text in pages:
        page = doc.new_page()
        y = 72
        for line in _wrap(text, 80):
            page.insert_text((72, y), line, fontsize=11)
            y += 14
            if y > 760:
                break
    if title or author:
        doc.set_metadata({"title": title, "author": author})
    doc.save(str(path))
    doc.close()
    return path


def _wrap(text: str, width: int) -> list[str]:
    lines, cur = [], ""
    for w in text.split():
        if len(cur) + len(w) + 1 > width:
            lines.append(cur); cur = w
        else:
            cur = (cur + " " + w).strip()
    if cur:
        lines.append(cur)
    return lines


def make_scanned_pdf(path: Path) -> Path:
    """A PDF whose pages are images without a text layer."""
    src = pymupdf.open()
    pg = src.new_page()
    pg.insert_text((72, 100), "scanned page content", fontsize=20)
    pix = pg.get_pixmap()
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_image(page.rect, pixmap=pix)
    doc.save(str(path))
    return path


def make_epub(path: Path, title: str, author: str, chapters: list[tuple[str, str]], lang: str = "en",
              year: str = "", isbn: str = "") -> Path:
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        z.writestr("META-INF/container.xml",
                   '<?xml version="1.0"?><container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
                   '<rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>')
        items = "".join(f'<item id="c{i}" href="c{i}.xhtml" media-type="application/xhtml+xml"/>' for i in range(len(chapters)))
        spine = "".join(f'<itemref idref="c{i}"/>' for i in range(len(chapters)))
        ident = f'<dc:identifier id="id">urn:isbn:{isbn}</dc:identifier>' if isbn else '<dc:identifier id="id">x</dc:identifier>'
        date = f"<dc:date>{year}</dc:date>" if year else ""
        z.writestr("OEBPS/content.opf",
                   '<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="id">'
                   f'<metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>{title}</dc:title><dc:creator>{author}</dc:creator>'
                   f'<dc:language>{lang}</dc:language>{date}{ident}</metadata>'
                   f'<manifest><item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>{items}</manifest>'
                   f'<spine>{spine}</spine></package>')
        lis = "".join(f'<li><a href="c{i}.xhtml">{t}</a></li>' for i, (t, _) in enumerate(chapters))
        z.writestr("OEBPS/nav.xhtml",
                   '<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops"><body>'
                   f'<nav epub:type="toc"><ol>{lis}</ol></nav></body></html>')
        for i, (t, body) in enumerate(chapters):
            paras = "".join(f"<p>{p}</p>" for p in body.split("\n\n"))
            z.writestr(f"OEBPS/c{i}.xhtml",
                       f'<html xmlns="http://www.w3.org/1999/xhtml"><head><title>{t}</title></head><body><h1>{t}</h1>{paras}</body></html>')
    return path


def make_docx(path: Path, title: str, paragraphs: list[str], author: str = "") -> Path:
    import docx
    d = docx.Document()
    d.core_properties.title = title
    d.core_properties.author = author
    d.add_heading(title, 0)
    for p in paragraphs:
        d.add_paragraph(p)
    d.save(str(path))
    return path


def palmdoc_compress_literal(data: bytes) -> bytes:
    """Valid (but unoptimised) PalmDOC stream: literals only."""
    out = bytearray()
    i = 0
    while i < len(data):
        b = data[i]
        if b == 0 or 9 <= b <= 0x7F:
            out.append(b)
            i += 1
        else:
            j = i
            while j < len(data) and j - i < 8 and not (data[j] == 0 or 9 <= data[j] <= 0x7F):
                j += 1
            out.append(j - i)
            out += data[i:j]
            i = j
    return bytes(out)


def make_mobi(path: Path, title: str, author: str, chapters: list[tuple[str, str]], compress: bool = True,
              encryption: int = 0, compression_override: int | None = None, cover_png: bytes | None = None) -> Path:
    import struct
    html = "<html><head></head><body>" + "".join(f"<h1>{t}</h1><p>{b}</p><mbp:pagebreak/>" for t, b in chapters) + "</body></html>"
    text = html.encode("utf-8")
    recs, pos = [], 0
    while pos < len(text):
        end = min(len(text), pos + 4096)
        while end < len(text) and (text[end] & 0xC0) == 0x80:    # do not split a UTF-8 character
            end -= 1
        recs.append(text[pos:end])
        pos = end
    comp = 2 if compress else 1
    body = [palmdoc_compress_literal(r) if compress else r for r in recs]
    if compression_override is not None:
        comp = compression_override
    images = [cover_png] if cover_png else []
    name = title.encode("utf-8")
    exth_items = [(100, author.encode("utf-8")), (101, b"Test Press"), (106, b"2011-05-01"), (524, b"ro")]
    if cover_png:
        exth_items.append((201, struct.pack(">I", 0)))
    exth_body = b"".join(struct.pack(">II", t, len(v) + 8) + v for t, v in exth_items)
    exth = b"EXTH" + struct.pack(">II", len(exth_body) + 12, len(exth_items)) + exth_body
    exth += b"\0" * (-len(exth) % 4)
    mh = bytearray(232)
    mh[0:4] = b"MOBI"
    struct.pack_into(">I", mh, 4, 232)
    struct.pack_into(">I", mh, 12, 65001)
    struct.pack_into(">I", mh, 20, 6)
    first_image = len(recs) + 1
    name_off = 16 + 232 + len(exth)
    struct.pack_into(">II", mh, 68, name_off, len(name))
    struct.pack_into(">I", mh, 92, first_image if images else 0xFFFFFFFF)
    struct.pack_into(">I", mh, 112, 0x40)
    r0 = struct.pack(">HHIHHHH", comp, 0, len(text), len(recs), 4096, encryption, 0) + bytes(mh) + exth + name + b"\0\0"
    records = [r0] + body + images
    n = len(records)
    header = bytearray(78 + 8 * n + 2)
    header[0:len(title[:31])] = title[:31].encode("ascii", "replace")
    header[60:68] = b"BOOKMOBI"
    struct.pack_into(">H", header, 76, n)
    off = len(header)
    for i, r in enumerate(records):
        struct.pack_into(">I", header, 78 + 8 * i, off)
        off += len(r)
    path.write_bytes(bytes(header) + b"".join(records))
    return path


def build_collection(root: Path) -> dict[str, dict]:
    """Create a nested synthetic library. Returns expectations keyed by relative path."""
    root.mkdir(parents=True, exist_ok=True)
    exp: dict[str, dict] = {}
    g = root / "Gardening" / "Permaculture"
    w = root / "Woodworking" / "Hand Tools"
    h = root / "Misc" / "Deep" / "Er"
    for d in (g, w, h):
        d.mkdir(parents=True, exist_ok=True)

    make_pdf(g / "food_forest.pdf", [
        "Introduction. This guide covers guilds. zanzibarium is a unique marker word.",
        "Page two discusses companion planting with tomatoes and basil. Tomatoes love basil.",
        "Page three: soil fertility without synthetic fertilizer through compost and mulch.",
    ], title="Food Forest Handbook", author="Ana Popescu")
    exp["Gardening/Permaculture/food_forest.pdf"] = {"title": "Food Forest Handbook", "authors": ["Ana Popescu"], "pages": 3}

    make_epub(g / "soil_book.epub", "Living Soil", "Ion Ionescu", [
        ("Preface", "A short preface about soil biology and quixotical earthworms."),
        ("Composting", "Composting turns waste into humus.\n\nHot composting needs a ratio of carbon to nitrogen."),
        ("Cover Crops", "Cover crops such as clover protect the ground. Mulch also helps."),
    ], lang="en", year="2015", isbn="9780306406157")
    exp["Gardening/Permaculture/soil_book.epub"] = {"title": "Living Soil", "authors": ["Ion Ionescu"], "chapters": 3}

    (w / "joints.txt").write_text(
        "Dovetail joints are strong. A beginner can cut a dovetail with a saw and chisel.\n\n"
        "The mortise and tenon is the classic frame joint. Zygomorphic is a unique marker.\n", encoding="utf-8")
    (w / "notes.md").write_text(
        "---\ntitle: Finishing Notes\nauthor: Maria Georgescu\n---\n# Finishing Notes\n\nShellac is a **traditional** finish. "
        "Use [oil](http://x) for tables. Marker: wolframite.\n", encoding="utf-8")
    (h / "page.html").write_text(
        "<html><head><title>Solar Heating</title><meta name='author' content='Radu Marin'>"
        "<script>var evil='scriptmarker';</script></head><body><h1>Solar Heating</h1>"
        "<p>A passive solar greenhouse stores heat in thermal mass. Marker: xylophagous.</p></body></html>", encoding="utf-8")
    make_docx(h / "energy.docx", "Battery Storage", [
        "Battery storage efficiency depends on round-trip losses.", "Lithium cells vs lead acid: marker quasimodo."], author="Dan Ilie")
    (h / "legacy.rtf").write_text(r"{\rtf1\ansi\deff0 {\fonttbl{\f0 Times;}}Folklore of the hearth. Marker: ostrogothic.\par}", encoding="utf-8")
    (root / "Misc" / "romanian.txt").write_text(
        "Grădina este frumoasă. Ședința de la București a avut loc în Țara Românească. Marker: ștefănescu.\n", encoding="utf-8")
    exp["Misc/romanian.txt"] = {}
    # malformed / unsupported
    (root / "Misc" / "broken.pdf").write_bytes(b"%PDF-1.4\nthis is not a real pdf")
    (root / "Misc" / "broken.epub").write_bytes(b"not a zip at all")
    (root / "Misc" / "book.djvu").write_bytes(b"AT&TFORM....")
    (root / "Misc" / "ignored.xyz").write_text("not a book")
    make_scanned_pdf(root / "Misc" / "scanned.pdf")
    return exp


def build_bulk(root: Path, n: int, seed: int = 1) -> int:
    """Generate ``n`` small text/EPUB/PDF documents spread over nested folders (for benchmarks)."""
    rng = random.Random(seed)
    root.mkdir(parents=True, exist_ok=True)
    for i in range(n):
        d = root / f"topic{i % 20}" / f"sub{i % 7}"
        d.mkdir(parents=True, exist_ok=True)
        kind = i % 4
        body = [filler(rng, 400) for _ in range(6)]
        if kind == 0:
            make_pdf(d / f"doc{i}.pdf", [filler(rng, 300) for _ in range(4)], title=f"Document {i}", author=f"Author {i % 50}")
        elif kind == 1:
            make_epub(d / f"doc{i}.epub", f"Book {i}", f"Author {i % 50}", [(f"Ch {j}", body[j]) for j in range(4)])
        else:
            (d / f"doc{i}.txt").write_text("\n\n".join(body), encoding="utf-8")
    return n
